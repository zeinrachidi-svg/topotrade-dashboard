#!/usr/bin/env python3
"""Regenerates data/team/<token>.json (personal dashboards) from the live index.html.
Usage: TEAM_TOKENS='{"Georges":"...",...}' GH_TOKEN=... python3 build_team_data.py"""
import json,re,os,base64,datetime,urllib.request,urllib.error
from bs4 import BeautifulSoup
REPO="zeinrachidi-svg/topotrade-dashboard"; GH=os.environ["GH_TOKEN"]; tokens=json.loads(os.environ["TEAM_TOKENS"])
def api(method,path,body=None,raw=False):
    req=urllib.request.Request(f"https://api.github.com/repos/{REPO}/contents/{path}",data=json.dumps(body).encode() if body else None,method=method,
      headers={"Authorization":f"Bearer {GH}","Accept":"application/vnd.github+json","User-Agent":"team-build"})
    try: return json.loads(urllib.request.urlopen(req).read())
    except urllib.error.HTTPError as e: return {"err":e.code}
def push(path,text):
    cur=api("GET",path); body={"message":"Refresh personal dashboard data","branch":"main","content":base64.b64encode(text.encode()).decode()}
    if cur.get("sha"): body["sha"]=cur["sha"]
    r=api("PUT",path,body); print(path,"ok" if "content" in r else r)
cur=api("GET","index.html"); 
if cur.get("content"): HTML=base64.b64decode(cur["content"]).decode()
else: HTML=urllib.request.urlopen(urllib.request.Request(cur["download_url"],headers={"Authorization":f"Bearer {GH}","User-Agent":"x"})).read().decode()

s=BeautifulSoup(HTML,"html.parser")
def sect(t):
    for h in s.find_all("h2"):
        if h.get_text().strip().startswith(t): return h
def after(h):
    for el in h.next_siblings:
        if getattr(el,"name",None)=="h2": break
        if getattr(el,"name",None): yield el
def tbl(h,i=0):
    ts=[e for e in after(h) if e.name=="table"]
    t=ts[i]
    rows=t.find_all("tr")
    cols=[c.get_text(" ",strip=True) for c in rows[0].find_all(["th","td"])]
    data=[[c.decode_contents().strip() for c in r.find_all("td")] for r in rows[1:]]
    return cols,data
def txt(x): return BeautifulSoup(x,"html.parser").get_text(" ",strip=True)

PEOPLE=[("Georges","Georges Labaky","Sales rep, Americas"),("Claude","Claude Nakhle","Sales rep, Asia / MEA"),
("Mireille","Mireille Baho","Sales rep, Europe / Australia, draft QA, eBay"),("Daniel","Daniel Llorente","Sales rep, France"),
("Zein","Zein Rachidi","CEO"),("Celine","Celine Fleyfel","Marketing"),("Yara","Yara Korh","Team member")]
def mine(first,owner): return first.lower() in txt(owner).lower()



sc=sect("Sales Team Comparison")
ytd_cols,ytd=tbl(sc,0); mtd_cols,mtd=tbl(sc,1)
crm_cols,crm=tbl(sect("CRM Pipeline by Rep"))
def row(rows,full): 
    for r in rows:
        if txt(r[0])==full: return [txt(c) for c in r]
# recs
recs={}
rh=sect("Recommended Actions per Rep")
cur=None;mode=None
for el in after(rh):
    if el.name=="h3": cur=el.get_text(strip=True); recs[cur]={"deal":[],"admin":[]}; continue
    if el.name=="p" and cur:
        t=el.get_text(strip=True)
        mode="deal" if t.startswith("Deal-Closing") else "admin" if t.startswith("Admin") else mode
    if el.name=="ol" and cur and mode: recs[cur][mode]=[li.decode_contents().strip() for li in el.find_all("li",recursive=False)]
# Celine recs
ch=sect("Recommended Actions for Celine")
celine_recs=[li.decode_contents().strip() for li in next(e for e in after(ch) if e.name=="ol").find_all("li",recursive=False)]
# tables
q_cols,q=tbl(sect("Open Offers"))
d_cols,drafts=tbl(sect("Pending Drafts"))
sh_cols,ship=tbl(sect("Ongoing Shipments"))
n_cols,nud=tbl(sect("Self-Serve Nudge"))
f_cols,fol=tbl(sect("Rep Actions Follow-Up"))
admin=[]
for t in ["Topotrade LLC FZ","Digital Realities SASU (France)","9494-1440","Group / cross"]:
    h=sect(t); c,rows=tbl(h)
    for r in rows: admin.append([t.split(" (")[0].replace("9494-1440","9494-1440 Québec")]+r)
a_cols=["Company"]+c
hot=[]
hh=sect("Hot Leads")
for card in hh.find_all_next("div",class_=re.compile("card")):
    if card.find_previous("h2") is not hh: break
    hot.append(card.decode_contents().strip())
# marketing KPIs
m=s.find(id="marketing")
mk=[]
grid=m.find(class_=re.compile("kpi"))
for c in grid.find_all(class_=re.compile("kpi|card"),recursive=False) if grid else []:
    mk.append(c.get_text(" | ",strip=True))
R=(dict(ytd_cols=ytd_cols,ytd=[[txt(c) for c in r] for r in ytd],mtd_cols=mtd_cols,mtd=[[txt(c) for c in r] for r in mtd],
 crm_cols=crm_cols,crm=[[txt(c) for c in r] for r in crm],recs=recs,celine_recs=celine_recs,q_cols=q_cols,q=q,d_cols=d_cols,drafts=drafts,
 sh_cols=sh_cols,ship=ship,n_cols=n_cols,nud=nud,f_cols=f_cols,fol=fol,a_cols=a_cols,admin=admin,hot=hot,mk=mk))

PEOPLE=[("Georges","Georges Labaky","Sales rep, Americas"),("Claude","Claude Nakhle","Sales rep, Asia / MEA"),
("Mireille","Mireille Baho","Sales rep, Europe / Australia, draft QA, eBay"),("Daniel","Daniel Llorente","Sales rep, France"),
("Zein","Zein Rachidi","CEO"),("Celine","Celine Fleyfel","Marketing"),("Yara","Yara Korh","Team member")]
def mine(first,x): return first.lower() in txt(x).lower()
def find(rows,full): 
    for r in rows:
        if r[0]==full: return r
_a=s.find(class_="asof")
asof=(_a.get_text(strip=True).replace("As of ","") if _a else str(datetime.date.today()))+" (daily run)"
out={}
_m={r[0]:r for r in R["mtd"]}
TEAM_COLS=["Rep","RevRec YTD","GPRec YTD","Tx YTD","RevRec MTD","GPRec MTD","Tx MTD"]
TEAM_ROWS=[[r[0],r[1],r[2],r[3]]+(_m[r[0]][1:4] if r[0] in _m else ["-","-","-"]) for r in R["ytd"]]

# ---------- extra finance tables (billings comparison, personal P&Ls) ----------
import csv as _csv
def _money(v,cur="$"): 
    return "-" if v is None else ("-" if v<0 else "")+cur+f"{abs(v):,.0f}"
def _pct(a,b): return "n/a" if not b else f"{(a-b)/abs(b)*100:+.0f}%"
def _prev(tok):
    cur=api("GET",f"data/team/{tok}.json")
    try: return json.loads(base64.b64decode(cur["content"]).decode())
    except Exception: return {}
def billings_table(path,asof):
    import re as _re
    rows=list(_csv.reader(open(path,encoding="utf-8")))
    h=rows[0]; iD=h.index("Date"); iI=[i for i,x in enumerate(h) if x.replace(" ","")=="Invoicenumber"][0]; iB=h.index("Billings excl Tax")
    mo={m:i+1 for i,m in enumerate(["january","february","march","april","may","june","july","august","september","october","november","december"])}
    def pd(x):
        x=x.strip(); m=_re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})$",x)
        if m: return datetime.date(int(m[3]),int(m[2]),int(m[1]))
        m=_re.match(r"(\d{1,2}) ([A-Za-z]+) (\d{4})$",x)
        if m: return datetime.date(int(m[3]),mo[m[2].lower()],int(m[1]))
    def num(x):
        x=_re.sub(r"[^\d.\-]","",x or ""); return float(x) if x not in("","-",".") else 0.0
    ents=[("FR","Digital Realities SASU (France)"),("AE","Topotrade LLC FZ (UAE)"),("CA","9494-1440 Quebec (Canada)")]
    a={e:dict(y26=0,y25=0,m26=0,m25=0) for e,_ in ents}
    for r in rows[1:]:
        if len(r)<=iB: continue
        d=pd(r[iD])
        if not d: continue
        inv=r[iI].strip(); e="CA" if inv.startswith("CA") else "AE" if inv.startswith("AE") else "FR"
        b=num(r[iB])
        for yr,k in((asof.year,"26"),(asof.year-1,"25")):
            cut=datetime.date(yr,asof.month,min(asof.day,28 if (asof.month==2 and asof.day>28) else asof.day))
            if d.year==yr and d<=cut:
                a[e]["y"+k]+=b
                if d.month==asof.month: a[e]["m"+k]+=b
    tot=dict(y26=0,y25=0,m26=0,m25=0); out=[]
    for e,name in ents:
        v=a[e]; [tot.__setitem__(k,tot[k]+v[k]) for k in v]
        out.append([name,_money(v["y26"],"EUR "),_money(v["y25"],"EUR "),_pct(v["y26"],v["y25"]),_money(v["m26"],"EUR "),_money(v["m25"],"EUR "),_pct(v["m26"],v["m25"])])
    out.append(["<b>Total</b>"]+[f"<b>{x}</b>" for x in [_money(tot["y26"],"EUR "),_money(tot["y25"],"EUR "),_pct(tot["y26"],tot["y25"]),_money(tot["m26"],"EUR "),_money(tot["m25"],"EUR "),_pct(tot["m26"],tot["m25"])]])
    mn=asof.strftime("%b")
    return dict(title="Billings by billing company: YTD and MTD vs last year",
        cols=["Billing company",f"YTD {asof.year}",f"YTD {asof.year-1}","YoY",f"MTD {mn} {asof.year}",f"MTD {mn} {asof.year-1}","YoY"],rows=out,
        note=f"Billings excl. tax in EUR from the Cockpit Sales tab, 1 Jan to {asof.day} {mn} each year. Entity = invoice prefix (CA = Canada, AE = UAE, numbers or blank = France, so not-yet-invoiced Topotrade sales are included). Live sheet, so it can differ slightly from the main dashboard snapshot.")

_SALES={}
def _prior_year(path,revcol,gpcol,asof,txcol=None,year=None):
    """Same-method monthly RevRec/GPRec for the previous year from the workbook Sales tab (day-capped for the current month)."""
    import openpyxl,re,datetime as _dt
    if path not in _SALES:
        _ws=openpyxl.load_workbook(path,data_only=True,read_only=True)["Sales"]
        _all=list(_ws.iter_rows(values_only=True)); _SALES[path]=(list(_all[0]),_all[1:])
    hdr,_rows=_SALES[path]; ix={}
    for i,h in enumerate(hdr):
        if h and h not in ix: ix[h]=i
    def num(v):
        if v is None: return 0.0
        if isinstance(v,(int,float)): return float(v)
        t=re.sub(r"[^\d.\-]","",str(v)); return float(t) if t not in("","-",".") else 0.0
    def pd(v):
        if isinstance(v,_dt.datetime): return v
        for f in("%d/%m/%Y","%d %B %Y","%Y-%m-%d"):
            try: return _dt.datetime.strptime(str(v).strip(),f)
            except Exception: pass
    out={m:[0.0,0.0,0.0] for m in range(1,13)}; y=year or asof.year-1
    for r in _rows:
        d=pd(r[ix["Date"]])
        if not d or d.year!=y: continue
        if y<asof.year and d.month==asof.month and d.day>asof.day: continue
        if y==asof.year and d>_dt.datetime(asof.year,asof.month,asof.day,23,59): continue
        a=out[d.month]; a[0]+=num(r[ix[revcol]]); a[1]+=num(r[ix[gpcol]]); a[2]+=num(r[ix[txcol]]) if txcol else 0
    return out
def _tx(v):
    return f"{v:g}" if v else "0"
def _yoy(cur,prev):
    if not prev: return "n/a"
    v=(cur/prev-1)*100
    return ">+999%" if v>999 else f"{v:+.0f}%"
def pl_canada(path):
    import openpyxl
    ws=openpyxl.load_workbook(path,data_only=True)["P&L Canada"]
    rows=[list(r) for r in ws.iter_rows(values_only=True)]
    st=next(i for i,r in enumerate(rows) if r[1] and "Revenue & GP recognition" in str(r[1]))
    out=[];T=dict(rev=0,gp=0,fc=0,np=0,tr=0,tg=0); PY=_prior_year(path,"RevRec GL","GPRec GL",_asof_d,"TXW GL"); CY=_prior_year(path,"RevRec GL","GPRec GL",_asof_d,"TXW GL",_asof_d.year); P25=[0,0,0]; C26=0
    for r in rows[st+2:]:
        if r[1]=="Total" or not hasattr(r[1],"year"): break
        if r[2] is None: continue
        out.append([r[1].strftime("%b %Y"),_money(r[2]),_money(r[6]),_money(r[7]),_money(r[8]),_money(r[9]),_money(r[10]),_money(r[11])]+([_money(PY[r[1].month][0]),_money(PY[r[1].month][1]),_yoy(r[2],PY[r[1].month][0]),_tx(CY[r[1].month][2]),_tx(PY[r[1].month][2])] if r[1].month in PY else ["","","","",""]))
        if r[1].month in PY: P25[0]+=PY[r[1].month][0]; P25[1]+=PY[r[1].month][1]; P25[2]+=PY[r[1].month][2]; C26+=CY[r[1].month][2]
        T["rev"]+=r[2];T["gp"]+=r[6];T["fc"]+=r[7];T["np"]+=r[8];T["tr"]+=r[10] or 0;T["tg"]+=r[11] or 0
    out.append(["<b>YTD</b>"]+[f"<b>{x}</b>" for x in [_money(T["rev"]),_money(T["gp"]),_money(T["fc"]),_money(T["np"]),"",_money(T["tr"]),_money(T["tg"]),_money(P25[0]),_money(P25[1]),_yoy(T["rev"],P25[0]),_tx(C26),_tx(P25[2])]])
    out.append(["<b>% of YTD target</b>",f"<b>{T['rev']/T['tr']*100:.0f}%</b>" if T["tr"] else "",f"<b>{T['gp']/T['tg']*100:.0f}%</b>" if T["tg"] else "","","","","","","","","",""])
    return dict(title="Topotrade Canada: YTD P&L (USD, revenue and GP recognition)",cols=["Month","RevRec","GPRec","Fixed cost","Net profit","Cumulated profit","RevRec target","GPRec target","RevRec 2025","GPRec 2025","RevRec YoY","Tx credited 2026","Tx credited 2025"],rows=out,
        note="From the 'Cockpit Report Canada' workbook, P&L Canada tab. YTD = months with actuals; later months only carry budgeted fixed costs. 2025 columns: same RevRec and GPRec method from the Sales tab, same months (current month to the same day). Tx credited = transactions counted with the same split-credit shares as RevRec (a shared deal counts 0.5), so it can differ from the deal counts in the team table. Net profit is not compared because 2025 fixed costs are held in a different currency basis.")
def pl_france(path):
    import openpyxl
    ws=openpyxl.load_workbook(path,data_only=True)["P&L"]
    rows=[list(r) for r in ws.iter_rows(values_only=True)]
    g=[];cur=None
    for c in rows[2]:
        cur=c or cur; g.append(cur)
    names=rows[3]; out=[];T=dict(rev=0,gp=0,fr=0,pl=0); grp={}; PY=_prior_year(path,"RevRec DL EUR","GPRec DL EUR",_asof_d,"TXW DL"); CY=_prior_year(path,"RevRec DL EUR","GPRec DL EUR",_asof_d,"TXW DL",_asof_d.year); P25=[0,0,0]; C26=0
    for r in rows[4:]:
        if not hasattr(r[1],"year") : break
        if r[2] is None or r[1].month>_asof_d.month: continue
        out.append([r[1].strftime("%b %Y")+(" (month to date)" if False else ""),_money(r[2],"EUR "),_money(r[3],"EUR "),_money(r[5],"EUR "),_money(r[4],"EUR ")]+([_money(PY[r[1].month][0],"EUR "),_money(PY[r[1].month][1],"EUR "),_yoy(r[2],PY[r[1].month][0]),_tx(CY[r[1].month][2]),_tx(PY[r[1].month][2])] if r[1].month in PY else ["","","","",""]))
        if r[1].month in PY: P25[0]+=PY[r[1].month][0]; P25[1]+=PY[r[1].month][1]; P25[2]+=PY[r[1].month][2]; C26+=CY[r[1].month][2]
        T["rev"]+=r[2];T["gp"]+=r[3];T["fr"]+=r[5] or 0;T["pl"]+=r[4] or 0
        for i in range(6,len(r)):
            if isinstance(r[i],(int,float)) and g[i]: grp[g[i]]=grp.get(g[i],0)+r[i]
    out.append(["<b>YTD</b>"]+[f"<b>{x}</b>" for x in [_money(T["rev"],"EUR "),_money(T["gp"],"EUR "),_money(T["fr"],"EUR "),_money(T["pl"],"EUR "),_money(P25[0],"EUR "),_money(P25[1],"EUR "),_yoy(T["rev"],P25[0]),_tx(C26),_tx(P25[2])]])
    t1=dict(title="Digital Realities France: YTD P&L (EUR)",cols=["Month","RevRec","GPRec","Total costs","P&L","RevRec 2025","GPRec 2025","RevRec YoY","Tx credited 2026","Tx credited 2025"],rows=out,
        note="From the 'Cockpit Report France DR' workbook, P&L tab. The latest month is month-to-date. Nov and Dec only hold budgeted costs and are not in YTD. 2025 columns: same RevRec and GPRec method from the Sales tab, same months (current month to the same day). Tx credited = transactions counted with the same split-credit shares as RevRec, so it can differ from the deal counts in the team table. Costs and P&L are not compared because the workbook holds no 2025 cost data.")
    t2=dict(title="Costs YTD by category (EUR)",cols=["Category","YTD"],rows=[[k,_money(v,"EUR ")] for k,v in sorted(grp.items(),key=lambda x:-x[1])])
    return [t1,t2]
_asof_d=datetime.date.today()
if os.environ.get("ASOF"): _asof_d=datetime.date.fromisoformat(os.environ["ASOF"])
EXTRA={}
if os.path.exists(os.environ.get("SALES_CSV","/nonexistent")):
    _b=billings_table(os.environ["SALES_CSV"],_asof_d); EXTRA["Georges"]=[_b]; EXTRA["Mireille"]=[_b]
if os.path.exists(os.environ.get("CANADA_XLSX","/nonexistent")): EXTRA.setdefault("Georges",[]).append(pl_canada(os.environ["CANADA_XLSX"]))
if os.path.exists(os.environ.get("FRANCE_XLSX","/nonexistent")): EXTRA["Daniel"]=pl_france(os.environ["FRANCE_XLSX"])
KEEP_TITLES=("Billings by billing company","Topotrade Canada: YTD P&L","Digital Realities France: YTD P&L","Costs YTD by category")


# ================= COACHING LAYER (today list, pace, attention, drafts, playbook, wins, board, manager) =================
import calendar as _cal
INIT={"Georges":"GL","Daniel":"DL","Mireille":"MB","Claude":"CN","Zein":"ZR"}
MONTHLY_TARGET={"Georges":86000}   # RevRec USD/month, from the Canada Cockpit. Others: self-benchmark until targets are confirmed.
def _num(x):
    x=re.sub(r"[^\d.\-]","",str(x or "")); 
    try: return float(x) if x not in("","-",".") else 0.0
    except Exception: return 0.0
def _task(P,start):
    for t in P["tasks"]:
        if t["title"].startswith(start): return t
    return None
def _rows(P,start):
    t=_task(P,start); return t["rows"] if t else []
def _wd_left(d):
    last=_cal.monthrange(d.year,d.month)[1]; n=0
    for day in range(d.day+1,last+1):
        if datetime.date(d.year,d.month,day).weekday()<5: n+=1
    return n
def coach_stats(P):
    ov=[txt(r[1])+" ("+txt(r[2])+")" for r in _rows(P,"Yesterday") if "overdue" in txt(r[3]).lower()]
    st=[r for r in _rows(P,"Self-serve") if (_num(r[5])>=7 or txt(r[7]).strip().upper().startswith("Y"))]
    dr=[r for r in _rows(P,"Drafts") if _num(r[3])>=2]
    sh=[r for r in _rows(P,"Shipments") if "BLOCKED" in txt(r[0]).upper()]
    oq=[r for r in _rows(P,"Open quotations") if _num(r[6])>7]
    risk=sum(_num(r[4]) for r in oq)
    return dict(overdue=ov,stale=st,drafts=dr,blocked=sh,oldq=oq,risk=risk)
def _short(x,n=190):
    x=txt(x); return x if len(x)<=n else x[:n-1]+"…"
def build_today(P,S):
    c=[]
    deal=P["recs"][0]["items"] if P["recs"] else []
    if deal: c.append(dict(text=_short(deal[0]),why="Top deal-closing action in today's recommendations."))
    for a in S["overdue"]: c.append(dict(text="Close the overdue action: "+a,why="Overdue items stay on this list until done."))
    for r in S["stale"]: c.append(dict(text=f"Call or write to {_short(r[0],70)} about {_short(r[2],60)}",why=f"Self-serve nudge pending {txt(r[5])} days"+(", escalated." if txt(r[7]).upper().startswith("Y") else ".")))
    for r in S["oldq"]: c.append(dict(text=f"Follow up quote {txt(r[0])} to {txt(r[1])} ({txt(r[4])})",why=f"Sent {txt(r[6])} days ago with no answer."))
    for r in S["blocked"]: c.append(dict(text=f"Unblock: {_short(r[3],80)} ({txt(r[1])})",why="Shipment is blocked; every day adds cost and customer risk."))
    for r in S["drafts"]: c.append(dict(text=f"Review the draft to {_short(r[1],70)}",why=f"Draft waiting {txt(r[3])} days."))
    for it in deal[1:]: c.append(dict(text=_short(it),why="Next deal-closing action."))
    return c[:3]
def build_pace(first,full,mtd_row,ytd_row,d):
    if not mtd_row or not ytd_row: return None
    mt=_num(mtd_row[1]); y=_num(ytd_row[1]); frac=(d.month-1+d.day/_cal.monthrange(d.year,d.month)[1])
    tgt=MONTHLY_TARGET.get(first); lab="monthly RevRec target"
    if not tgt: tgt=y/frac if frac else 0; lab="your YTD monthly average (no target set yet)"
    left=_wd_left(d); gap=max(tgt-mt,0); dayfrac=d.day/_cal.monthrange(d.year,d.month)[1]
    return dict(mtd=_money(mt),target=_money(tgt),label=lab,gap=_money(gap),daysLeft=left,perDay=_money(gap/left) if left else _money(gap),
        pct=round(mt/tgt*100) if tgt else 0,expected=round(dayfrac*100),month=d.strftime("%B"))
def email_for(r):
    m=re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+",txt(r[0])); return m.group(0) if m else ""
def build_drafts(first,S,P):
    out=[]
    rows=S["stale"]+[r for r in _rows(P,"Self-serve") if r not in S["stale"]]
    for r in rows[:4]:
        lead=re.sub(r"\s*\(.*","",txt(r[0])).strip() or "there"; item=txt(r[2]); intent=txt(r[1]).lower()
        if intent.startswith("sell"):
            body=f"Hi,\n\nFollowing up on your {item}. We have active buyers and open RFPs for this type of equipment, so I would like to list it this week.\nCould you send the serial numbers, condition, a few photos and your target price?\n\nBest regards,\n{first}"
            sub=f"Listing your {item}: next step"
            wa=f"Hi, it's {first} from Topotrade. We have buyers for your {item}. Can you send serial numbers, photos and your target price so I can list it this week?"
        else:
            body=f"Hi,\n\nFollowing up on your interest in {item}. I can confirm availability and send you a quote today. Is this still on your radar, and what timeline are you working to?\n\nBest regards,\n{first}"
            sub=f"{item}: availability and quote"
            wa=f"Hi, it's {first} from Topotrade. Still interested in {item}? I can confirm availability and send a quote today."
        out.append(dict(title=f"Follow-up: {lead}",to=email_for(r),subject=sub,body=body,wa=wa))
    for r in S["oldq"][:2]:
        body=f"Hi,\n\nI am following up on our quotation {txt(r[0])} for {txt(r[2])} ({txt(r[4])}). Do you have any questions, or can I help you move forward this week?\n\nBest regards,\n{first}"
        out.append(dict(title=f"Quote follow-up: {txt(r[1])}",to="",subject=f"Quotation {txt(r[0])}: following up",body=body,wa=f"Hi, it's {first} from Topotrade. Any questions on quotation {txt(r[0])}? Happy to adjust to move forward."))
    return out
def build_playbook(first,S,P):
    has=lambda s: _task(P,s) is not None
    R_=[]
    R_.append(("No overdue actions",not S["overdue"],"auto"))
    if has("Self-serve"): R_.append(("Self-serve nudges followed up within 7 days",not S["stale"],"auto"))
    if has("Drafts"): R_.append(("Drafts reviewed within 48 hours",not S["drafts"],"auto"))
    if has("Open quotations"): R_.append(("Open quotations followed up within 7 days",not S["oldq"],"auto"))
    if has("Shipments"): R_.append(("Blocked shipments have a next step today",not S["blocked"],"auto"))
    for s in ["Reply to every new lead within 24 hours","Log a note in Odoo after every call or meeting","Contact at least 3 new marketplace users or sellers this week"]: R_.append((s,None,"self"))
    return [dict(rule=a,ok=b,kind=k) for a,b,k in R_]
def sales_rows(path):
    rows=list(_csv.reader(open(path,encoding="utf-8"))); h=rows[0]; return h,rows[1:]
def _pd(x):
    x=x.strip(); m=re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})$",x)
    mo={m_:i+1 for i,m_ in enumerate(["january","february","march","april","may","june","july","august","september","october","november","december"])}
    if m: return datetime.date(int(m[3]),int(m[2]),int(m[1]))
    m=re.match(r"(\d{1,2}) ([A-Za-z]+) (\d{4})$",x)
    if m and m[2].lower() in mo: return datetime.date(int(m[3]),mo[m[2].lower()],int(m[1]))
def wins_for(first,path,d):
    h,rows=sales_rows(path); ini=INIT.get(first); out=[]; tot=0
    for r in rows:
        if len(r)<19: continue
        dt=_pd(r[4])
        if not dt or not(d-datetime.timedelta(days=6)<=dt<=d): continue
        if (r[5].strip().upper()!=ini): continue
        b=_num(r[16]); tot+=b
        out.append([dt.strftime("%d %b"),r[14] or "-",r[7] or r[2] or "-",_money(b,"EUR "),_money(_num(r[18]),"EUR ")])
    return out,tot
def team_week(path,d):
    h,rows=sales_rows(path); res={}
    for f,ini in INIT.items():
        n=0;tot=0
        for r in rows:
            if len(r)<19: continue
            dt=_pd(r[4])
            if dt and d-datetime.timedelta(days=6)<=dt<=d and r[5].strip().upper()==ini: n+=1; tot+=_num(r[16])
        res[f]=(n,tot)
    return res
HIST_PATH="data/team/history.json"
try: HIST=json.loads(base64.b64decode(api("GET",HIST_PATH)["content"]).decode())
except Exception: HIST={}
_dkey=_asof_d.isoformat()
WEEK=team_week(os.environ["SALES_CSV"],_asof_d) if os.path.exists(os.environ.get("SALES_CSV","/nonexistent")) else {}
PENDING={}; STATS={}
for first,full,role in PEOPLE:
    P=dict(team=dict(title="Team comparison",cols=TEAM_COLS,rows=TEAM_ROWS,hl=full),name=full,first=first,role=role,asOf=asof,updatedAt=datetime.datetime.utcnow().isoformat()+"Z",kpis=[],tables=[],recs=[],tasks=[],note="")
    if first in("Georges","Claude","Mireille","Daniel","Zein"):
        y=find(R["ytd"],full); mt=find(R["mtd"],full); team=find(R["ytd"],"Team"); tmt=find(R["mtd"],"Team"); c=find(R["crm"],full)
        P["kpis"]=[dict(label="Revenue YTD",value=y[1],sub=f"{y[6]} vs same period 2025"),dict(label="Gross profit YTD",value=y[2],sub=f"GM {y[4]}"),
          dict(label="Transactions YTD",value=y[3],sub=f"team total {team[3]}"),dict(label="Revenue MTD (Oct)",value=mt[1],sub=f"{mt[3]} tx, GP {mt[2]}"),
          dict(label="Open opportunities",value=c[1],sub=f"expected {c[2]}")]
        P["tables"]=[dict(title="Year to date vs last year",cols=R["ytd_cols"],rows=R["ytd"],hl=full),
                     dict(title="Month to date",cols=R["mtd_cols"],rows=R["mtd"],hl=full),
                     dict(title="Your biggest open deals",cols=["Deal"],rows=[[d.strip()] for d in c[3].split("; ")])]
        P["recs"]=[dict(title="Deal-closing actions",items=R["recs"][full]["deal"]),dict(title="Admin & housekeeping",items=R["recs"][full]["admin"])]
        q=[r for r in R["q"] if txt(r[2])==full]
        if first=="Zein": q+= [r for r in R["q"] if txt(r[2]) in("Contact","(none)")]
        P["tasks"].append(dict(title="Open quotations to follow up"+(" (incl. unassigned)" if first=="Zein" else ""),cols=R["q_cols"],rows=q,
            note="All are 5+ business days old with no confirmation. Pick the recent, real-customer ones first (see your deal-closing list)."))
        pk=[r for r in R["ship"] if mine(first,r[4])]
        P["tasks"].append(dict(title="Shipments you own",cols=R["sh_cols"],rows=pk))
        P["tasks"].append(dict(title="Self-serve nudges you own",cols=R["n_cols"],rows=[r for r in R["nud"] if mine(first,r[3])]))
        P["tasks"].append(dict(title="Admin tasks you own",cols=R["a_cols"],rows=[r for r in R["admin"] if mine(first,r[5])]))
        P["tasks"].append(dict(title="Yesterday's actions, follow-up",cols=R["f_cols"],rows=[r for r in R["fol"] if mine(first,r[0])]))
        dr=R["drafts"] if first=="Mireille" else []
        P["tasks"].append(dict(title="Drafts waiting for review"+("" if first=="Mireille" else " (Mireille reviews all drafts)"),cols=R["d_cols"],rows=dr))
        hl=[h for h in R["hot"] if first.lower() in txt(h).lower()]
        P["tasks"].append(dict(title="Hot leads naming you",cols=["Lead"],rows=[[h] for h in hl]))
        if first=="Zein":
            P["tasks"][2]=dict(title="Self-serve nudges (all reps, company-wide)",cols=R["n_cols"],rows=R["nud"])
            P["tasks"][5]=dict(title="Drafts pending review (company-wide)",cols=R["d_cols"],rows=R["drafts"])
            P["tables"].insert(0,dict(title="Pipeline by rep (company)",cols=R["crm_cols"],rows=R["crm"]))
            P["tasks"][3]["title"]="Admin tasks you own (all companies)"
    elif first=="Celine":
        P["kpis"]=[dict(label=k.split(" | ")[0],value=k.split(" | ")[1].replace("●","").strip() or "n/a",sub=" ".join(k.split(" | ")[2:])) for k in R["mk"]]
        P["recs"]=[dict(title="Your 10 recommended actions",items=R["celine_recs"])]
        P["tasks"].append(dict(title="Drafts waiting for your approval",cols=R["d_cols"],rows=[r for r in R["drafts"] if re.search("campaign|social|LinkedIn|Facebook|post",txt(r[0]),re.I)]))
        P["tasks"].append(dict(title="Admin tasks you own",cols=R["a_cols"],rows=[r for r in R["admin"] if mine(first,r[5])]))
        P["note"]="Ready-to-use suggested posts (sold items, latest listings, wanted items) are on the main dashboard's Marketing tab."
    else:
        P["note"]="Your page is ready, but no daily data feeds it yet (no Odoo salesperson record or Cockpit sales rows under your name). Tell Zein if you should get a rep-style or marketing-style view."
    _need={"Georges":["Billings by billing company","Topotrade Canada: YTD P&L"],"Mireille":["Billings by billing company"],"Daniel":["Digital Realities France: YTD P&L","Costs YTD by category"]}.get(first,[])
    P["tables"]+=EXTRA.get(first,[])
    _have=[t["title"] for t in P["tables"]]
    if any(not any(h.startswith(n) for h in _have) for n in _need):
        for t in _prev(tokens[first]).get("tables",[]):
            if t["title"].startswith(KEEP_TITLES) and not any(h==t["title"] for h in _have): P["tables"].append(t); print("kept previous table for",first,":",t["title"])
    # ---- coaching enrichment ----
    S=coach_stats(P); STATS[first]=S; prevJ=_prev(tokens[first])
    P["today"]=build_today(P,S)
    if first in INIT:
        P["pace"]=build_pace(first,full,find(R["mtd"],full),find(R["ytd"],full),_asof_d)
        if WEEK:
            P["wins"]=dict(rows=wins_for(first,os.environ["SALES_CSV"],_asof_d)[0],total=_money(wins_for(first,os.environ["SALES_CSV"],_asof_d)[1],"EUR "))
        elif prevJ.get("wins"): P["wins"]=prevJ["wins"]
    P["attn"]=[dict(label="Overdue actions",n=len(S["overdue"]),detail="; ".join(S["overdue"])),
               dict(label="Stale nudges (7+ days)",n=len(S["stale"]),detail="; ".join(_short(r[0],50) for r in S["stale"])),
               dict(label="Drafts waiting 2+ days",n=len(S["drafts"]),detail="; ".join(_short(r[1],50) for r in S["drafts"])),
               dict(label="Blocked shipments",n=len(S["blocked"]),detail="; ".join(_short(r[3],50) for r in S["blocked"])),
               dict(label="Quotes unanswered 7+ days",n=len(S["oldq"]),detail=("at risk "+_money(S["risk"])) if S["oldq"] else "")]
    P["drafts"]=build_drafts(first,S,P) if first in INIT or first=="Zein" else []
    P["playbook"]=build_playbook(first,S,P)
    HIST.setdefault(_dkey,{})[first]=dict(overdue=len(S["overdue"]),stale=len(S["stale"]),drafts=len(S["drafts"]))
    PENDING[first]=P
    out[first]=sum(len(t["rows"]) for t in P["tasks"])
print(out)
# ---- history trim + streaks + accountability board + manager view ----
for k in sorted(HIST)[:-60]: HIST.pop(k)
def streak(f):
    n=0
    for k in sorted(HIST,reverse=True):
        h=HIST[k].get(f)
        if h is None: break
        if h["overdue"]==0 and h["stale"]==0: n+=1
        else: break
    return n
REPS=["Georges","Claude","Mireille","Daniel"]
board=[]
for f in REPS:
    S=STATS[f]; n,tot=WEEK.get(f,(None,None))
    board.append([f,str(len(S["overdue"])),str(len(S["stale"])),str(len(S["drafts"])),str(len(S["blocked"])),
        "-" if n is None else str(n),"-" if tot is None else _money(tot,"EUR "),f"{streak(f)} day"+("" if streak(f)==1 else "s")])
BOARD=dict(title="Team accountability board",cols=["Rep","Overdue","Stale nudges","Old drafts","Blocked ship.","Deals 7d","Billings 7d","Clean streak"],rows=board,
    note="Clean streak = consecutive daily runs with no overdue action and no stale nudge. It starts counting from the first run after this feature was added.")
mgr_rows=[]
for f in REPS:
    S=STATS[f]; pc=PENDING[f].get("pace") or {}
    mgr_rows.append([f,"<br>".join(S["overdue"]) or "-","<br>".join(_short(r[0],60)+" ("+txt(r[5])+"d)" for r in S["stale"]) or "-","<br>".join(_short(r[3],60) for r in S["blocked"]) or "-",
        (f"{pc['pct']}% of target, needs {pc['perDay']}/day" if pc else "-")])
MGR=dict(title="Manager view: who needs a push",cols=["Rep","Overdue actions","Stale nudges","Blocked shipments","Pace this month"],rows=mgr_rows)
for first,P in PENDING.items():
    P["board"]=BOARD
    if first=="Zein": P["manager"]=MGR
    push(f"data/team/{tokens[first]}.json",json.dumps(P,ensure_ascii=False))
push(HIST_PATH,json.dumps(HIST))

# ---------- ensure Team Dashboards tab exists in index.html ----------
def ensure_team_tab():
    import re as _re
    h=api("GET","index.html"); html=base64.b64decode(h["content"]).decode()
    if "<!--TEAMTAB-->" in html: print("team tab present"); return
    g=lambda p: base64.b64decode(api("GET",p)["content"]).decode()
    btn=g("scripts/team_tab_button.html"); blk=g("scripts/team_tab_block.html")
    a="<button class='tab-btn' onclick=\"showTab('automation')\">Automation Map</button>"
    t="<script>\nfunction showTab"
    if html.count(a)!=1 or html.count(t)!=1: print("team tab NOT restored: anchors missing"); return
    html=html.replace(a,a+btn).replace(t,blk+t)
    api("PUT","index.html",{"message":"Restore Team Dashboards tab","content":base64.b64encode(html.encode()).decode(),"sha":h["sha"],"branch":"main"}); print("team tab restored")

ensure_team_tab()

def ensure_kpis():
    """Add Net profit YTD/MTD and Transactions MTD cards to the main dashboard Top KPIs (daily run rebuilds index.html, so this re-applies)."""
    import openpyxl,re as _re,datetime as _dt
    path=os.environ.get("MASTER_XLSX","")
    if not os.path.exists(path): print("kpis skipped: no MASTER_XLSX"); return
    asof=_asof_d; wb=openpyxl.load_workbook(path,data_only=True)
    P=wb["P&L"]; r0=None
    for r in range(1,P.max_row+1):
        if str(P.cell(r,3).value).strip()==f"P&L {asof.year}": r0=r; break
    if not r0: print("kpis skipped: P&L block not found"); return
    net_ytd=0.0; cur=None
    for r in range(r0+1,r0+14):
        lab=P.cell(r,3).value
        if not hasattr(lab,"month") and not (isinstance(lab,str) and lab[:3].isalpha() and lab.strip()[-4:]==str(asof.year)): continue
        m=lab.month if hasattr(lab,"month") else _dt.datetime.strptime(lab.strip(),"%b %Y").month
        if m>asof.month: continue
        n=P.cell(r,11).value
        if isinstance(n,(int,float)) and P.cell(r,4).value is not None: net_ytd+=n
        if m==asof.month: cur=dict(net=n,gp2=P.cell(r,8).value,fix=P.cell(r,10).value)
    S=wb["Sales"]; hdr=[c.value for c in S[1]]; ix={}
    for i,h in enumerate(hdr):
        if h and h not in ix: ix[h]=i
    def pd(v):
        if isinstance(v,_dt.datetime): return v
        for f in("%d/%m/%Y","%d %B %Y"):
            try: return _dt.datetime.strptime(str(v).strip(),f)
            except Exception: pass
    ty=tm=0
    for r in S.iter_rows(min_row=2,values_only=True):
        d=pd(r[ix["Date"]])
        if not d or d.date()>asof or d.year!=asof.year: continue
        if not (r[ix["Product"]] or r[ix["Invoice  number"]]): continue
        ty+=1
        if d.month==asof.month: tm+=1
    eur=lambda v: ("-" if v<0 else "")+"&euro;"+f"{abs(v):,.0f}"
    def card(label,val,sub,neg=False):
        st=" style='color:#c0392b'" if neg else ""
        return f"<div class='kpi' data-kpi='plx'><div class='label'>{label}</div><div class='value'{st}>{val}</div><div class='sub'>{sub}</div></div>"
    mon=asof.strftime("%b")
    c_net=card("Net Profit YTD",eur(net_ytd),f"P&amp;L tab, EUR · Jan to {mon} {asof.year} · as of {asof.strftime('%d %b')}",net_ytd<0)
    c_mtd=card("Net Profit MTD",eur(cur["net"]) if cur and isinstance(cur["net"],(int,float)) else "n/a",
        (f"GP2 {eur(cur['gp2'])} less fixed costs {eur(cur['fix'])} (full month of costs against GP so far)" if cur and isinstance(cur.get('gp2'),(int,float)) else ""),bool(cur and isinstance(cur["net"],(int,float)) and cur["net"]<0))
    c_tx=card("Transactions MTD",str(tm),f"Cockpit rows dated {mon} {asof.year} · {ty} YTD")
    h=api("GET","index.html"); html=base64.b64decode(h["content"]).decode()
    html=_re.sub(r"<div class='kpi' data-kpi='plx'>.*?</div></div>","",html,flags=_re.S)
    a=_re.search(r"<div class='kpi'><div class='label'>Gross Profit YTD</div>.*?</div></div>",html,_re.S)
    b=_re.search(r"<div class='kpi'><div class='label'>Transactions YTD</div>.*?</div></div>",html,_re.S)
    if not a or not b: print("kpis NOT added: anchors missing"); return
    # insert later anchor first so offsets stay valid
    first,second=(a,b) if a.end()<b.end() else (b,a)
    ins={a.end():c_net+c_mtd, b.end():c_tx}
    for pos in sorted(ins,reverse=True): html=html[:pos]+ins[pos]+html[pos:]
    api("PUT","index.html",{"message":"Top KPIs: net profit YTD/MTD, transactions MTD","content":base64.b64encode(html.encode()).decode(),"sha":h["sha"],"branch":"main"}); print("kpis added",eur(net_ytd),tm,ty)

ensure_kpis()
